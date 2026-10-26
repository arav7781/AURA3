"""Sandbox hardening and office conversion. Needs the Docker daemon and the aura-sandbox image."""

import io
import json
import shutil
import time
import uuid
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from aura.sandbox.office import ConversionError, check_zip_bomb, convert_to_pdf
from aura.sandbox.runner import LABEL, RunSpec, runner

pytestmark = pytest.mark.docker


@pytest.fixture(scope="module", autouse=True)
def sandbox_ready():
    status = runner.status()
    if status != "ready":
        pytest.skip(f"sandbox not available ({status})")


@pytest.fixture
def workspace():
    # Colima shares only $HOME with the Docker VM, so test workspaces live there (not in /tmp).
    path = Path.home() / ".cache" / "aura-tests" / uuid.uuid4().hex
    path.mkdir(parents=True)
    yield path
    shutil.rmtree(path, ignore_errors=True)


PROBE = r"""
import json, os, socket
result = {}
try:
    socket.create_connection(("1.1.1.1", 80), timeout=2)
    result["network"] = "open"
except OSError as exc:
    result["network"] = "blocked"
try:
    open("/etc/aura-probe", "w").write("x")
    result["rootfs"] = "writable"
except OSError:
    result["rootfs"] = "read-only"
open("/work/out/ok.txt", "w").write("hello")
open("/tmp/scratch.txt", "w").write("tmp")
result["uid"] = os.getuid()
print(json.dumps(result))
"""


async def test_container_is_isolated_but_can_write_its_workspace(workspace):
    result = await runner.run(
        RunSpec(argv=["python", "/opt/aura/runner.py", "/work/main.py"], workspace=workspace,
                files={"main.py": PROBE.encode()}, timeout_s=60)
    )
    assert result.status == "ok", result.stderr
    probe = json.loads(result.stdout[-1])
    assert probe["network"] == "blocked"
    assert probe["rootfs"] == "read-only"
    assert (workspace / "out" / "ok.txt").read_text() == "hello"


async def test_runaway_code_is_killed_at_the_deadline(workspace):
    started = time.monotonic()
    result = await runner.run(
        RunSpec(argv=["python", "/opt/aura/runner.py", "/work/main.py"], workspace=workspace,
                files={"main.py": b"while True:\n    pass\n"}, timeout_s=3)
    )
    assert result.status == "timeout"
    assert time.monotonic() - started < 20
    assert runner.client().containers.list(all=True, filters={"label": LABEL}) == []


async def test_output_is_capped(workspace):
    result = await runner.run(
        RunSpec(argv=["python", "/opt/aura/runner.py", "/work/main.py"], workspace=workspace,
                files={"main.py": b"for i in range(5000):\n    print(i)\n"}, timeout_s=60)
    )
    assert result.status == "ok" and result.truncated and len(result.stdout) <= 2000


def _deck_with_hidden_slide(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for i, title in enumerate(["Problem", "Hidden appendix", "Ask"], start=1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2))
        box.text_frame.text = f"Slide {i}: {title}"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        if title.startswith("Hidden"):
            slide._element.set("show", "0")
    deck.save(path)


async def test_pptx_converts_with_hidden_slides_kept(workspace):
    src = workspace / "deck.pptx"
    _deck_with_hidden_slide(src)
    pdf = await convert_to_pdf(src, workdir=workspace / "convert")
    document = pdfium.PdfDocument(pdf)
    assert len(document) == 3
    texts = [document[i].get_textpage().get_text_range() for i in range(3)]
    assert "Hidden appendix" in texts[1]


def test_zip_bombs_are_rejected_before_conversion(tmp_path):
    bomb = tmp_path / "bomb.pptx"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/huge.xml", b"0" * (400 << 20))
    bomb.write_bytes(buffer.getvalue())
    with pytest.raises(ConversionError):
        check_zip_bomb(bomb)
