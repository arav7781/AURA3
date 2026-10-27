"""Seed decks: byte-stable builds, gold slide map, scanned and adversarial variants, registry fixtures."""

import io
from pathlib import Path

import pdfplumber
import pytest
from pptx import Presentation

from aura.decks.build import build_deck
from aura.decks.spec import load_spec, load_specs
from aura.evaluation import mca
from aura.ingest.injection import hidden_pdf_text, pptx_notes

SPECS = Path(__file__).parents[2] / "seed" / "specs"


def spec(slug: str):
    return next(load_spec(p) for p in SPECS.glob(f"*-{slug}.yaml"))


def page_text(pdf: bytes, page: int) -> str:
    """A page's text in drawing order with whitespace collapsed (slides lay text out in side-by-side cards,
    which position-based extraction interleaves)."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(pdf)
    return " ".join(document[page - 1].get_textpage().get_text_range().split())


def test_every_spec_is_valid_with_a_unique_slug():
    specs = load_specs(SPECS)
    assert len(specs) == 25 and len({s.slug for s in specs}) == 25


def test_builds_are_byte_stable_and_the_slide_map_holds_gold_pages():
    first, second = build_deck(spec("krishimitra-ai")), build_deck(spec("krishimitra-ai"))
    assert first.pdf == second.pdf and first.pptx == second.pptx
    facts = first.slide_map["facts"]
    assert facts["market.tam"] == {"page": 4, "text": "USD 82.3M in 2025"}
    assert facts["traction.arr"] == {"page": 6, "text": "ARR: Rs 3.1 crore in FY2025"}
    with pdfplumber.open(io.BytesIO(first.pdf)) as pdf:
        assert len(pdf.pages) == 11
    assert "USD 82.3M in 2025" in page_text(first.pdf, 4)
    assert len(Presentation(io.BytesIO(first.pptx)).slides) == 11


def test_scanned_decks_have_no_text_layer():
    files = build_deck(spec("loomlane"))
    assert files.slide_map["scanned"]
    with pdfplumber.open(io.BytesIO(files.pdf)) as pdf:
        assert len(pdf.pages) == 11 and not any(page.chars for page in pdf.pages)


def test_the_injection_deck_hides_its_instruction_and_others_do_not(tmp_path):
    attack = build_deck(spec("neocredit-score"))
    (tmp_path / "attack.pdf").write_bytes(attack.pdf)
    (tmp_path / "attack.pptx").write_bytes(attack.pptx)
    (tmp_path / "clean.pdf").write_bytes(build_deck(spec("krishimitra-ai")).pdf)
    assert [f.page for f in hidden_pdf_text(tmp_path / "attack.pdf")] == [7] == [attack.slide_map["injection_page"]]
    assert {f.page for f in pptx_notes(tmp_path / "attack.pptx")} == {7}
    assert hidden_pdf_text(tmp_path / "clean.pdf") == []


def test_an_inconsistent_deck_contradicts_itself_on_the_traction_slide():
    text = page_text(build_deck(spec("agrisense-networks")).pdf, 6)
    assert "ARR of Rs 1.4 crore across 3,200 deployed sensors" in text and "ARR: Rs 0.4 crore in FY2025" in text


@pytest.fixture
def registry():
    mca.registry_fixtures.clear()
    mca.load_registry_fixtures(SPECS)
    yield
    mca.registry_fixtures.clear()


@pytest.mark.parametrize(("slug", "flags"), [
    ("krishimitra-ai", []),
    ("portpulse", ["COMPANY_STATUS_STRUCK_OFF"]),
    ("suraksha-insure", ["NAME_MISMATCH"]),
    ("tattva-foods", ["INVALID_CIN_FORMAT"]),
])
def test_registry_fixtures_come_from_the_specs(app_env, registry, slug, flags):
    s = spec(slug)
    verification = mca.verify(s.mca.cin, legal_name=s.legal_name, claimed_directors=[m.name for m in s.team])
    assert mca.consistency_flags(verification, legal_name=s.legal_name) == flags
