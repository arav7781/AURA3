"""Prompt-injection scanning for uploaded documents.

Deck text is data, never instructions: agents always receive it wrapped as untrusted content. This scanner
adds a visible signal on top: it looks for hidden text (white or microscopic characters in the PDF),
instruction-like phrasing aimed at an AI, and the same in PPTX speaker notes. Hits flag the chunk
(``injection_suspected``) and the document (``DECK_HIDDEN_INSTRUCTIONS``) so reviewers and the UI see it.
"""

import re
from dataclasses import dataclass
from pathlib import Path

PATTERNS = [
    r"ignore (all |any |the )?(previous|prior|above|earlier|preceding) (instructions|prompts|directions|rules)",
    r"disregard (all |any |the )?(previous|prior|above|earlier) ",
    r"(note|message|instruction|instructions) (to|for) (the |any )?(ai|llm|model|evaluator|assistant|reviewer bot)",
    r"(you are|act as|pretend to be) (an? |the )?(ai|assistant|language model|evaluator|grader|analyst bot)",
    r"(rate|score|grade|rank) (this|the|our) (startup|company|deck|pitch|team)( as| at)? (10|ten|9\.?\d?)( ?/ ?10| out of (10|ten))?",
    r"(must|should) (receive|get|be given) (a )?(score|rating) of",
    r"system prompt|developer message|<\|im_start\|>|\[/?inst\]|### ?instruction",
    r"do not (mention|reveal|disclose|flag) (this|these|the following)",
    r"respond only with|output only",
]
_PATTERN = re.compile("|".join(f"(?:{p})" for p in PATTERNS), re.IGNORECASE)


@dataclass(frozen=True)
class Finding:
    page: int | None
    source: str  # hidden_text | text | notes
    excerpt: str


def looks_like_instruction(text: str) -> bool:
    return bool(_PATTERN.search(text))


def scan_text(text: str, page: int | None = None, source: str = "text") -> list[Finding]:
    found = []
    for match in _PATTERN.finditer(text):
        start = max(0, match.start() - 60)
        found.append(Finding(page, source, text[start : match.end() + 80].strip()))
    return found


def _is_white(color: object) -> bool:
    if color is None:
        return False
    values = color if isinstance(color, list | tuple) else [color]
    try:
        numbers = [float(v) for v in values]
    except (TypeError, ValueError):
        return False
    if len(numbers) == 4:  # CMYK white is (0, 0, 0, 0)
        return all(v <= 0.02 for v in numbers)
    return bool(numbers) and all(v >= 0.97 for v in numbers)


def hidden_pdf_text(pdf_path: Path) -> list[Finding]:
    """Text a human cannot see: white fill or a font smaller than 2 pt."""
    import pdfplumber

    findings: list[Finding] = []
    with pdfplumber.open(pdf_path) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            hidden = [
                ch["text"]
                for ch in page.chars
                if _is_white(ch.get("non_stroking_color")) or float(ch.get("size") or 12) < 2.0
            ]
            text = "".join(hidden).strip()
            if len(text) >= 12:
                findings.append(Finding(index, "hidden_text", text[:400]))
    return findings


def pptx_notes(pptx_path: Path) -> list[Finding]:
    from pptx import Presentation

    findings: list[Finding] = []
    deck = Presentation(str(pptx_path))
    for index, slide in enumerate(deck.slides, start=1):
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            findings += scan_text(notes, index, "notes")
    return findings


def scan_document(pdf_path: Path | None, *, original: Path | None = None) -> list[Finding]:
    findings: list[Finding] = []
    if pdf_path is not None and pdf_path.suffix.lower() == ".pdf":
        findings += hidden_pdf_text(pdf_path)
    if original is not None and original.suffix.lower() == ".pptx":
        findings += pptx_notes(original)
    return findings
