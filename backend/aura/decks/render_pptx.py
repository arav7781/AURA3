"""Slides → a PowerPoint deck drawn from the same ``design.lay_out`` shapes as the PDF (landscape A4, so every
position maps one to one). Text uses the Geist family; where it is not installed PowerPoint substitutes its
default sans. Output is byte-stable: core properties carry a fixed date and every zip entry a fixed timestamp."""

import io
import zipfile
from datetime import datetime

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from aura.decks.design import Brand, Disc, Fill, H, Line, Page, Picture, Shadow, W, lay_out
from aura.decks.slides import Slide, build_slides
from aura.decks.spec import StartupSpec
from aura.decks.theme import JADE_900, PPTX_FONTS, RGB, WASH

FIXED_DATE = datetime(2026, 1, 1)
EMU = 12700  # per point


def _emu(points: float) -> Emu:
    return Emu(round(points * EMU))


def _hex(rgb: RGB) -> str:
    return "{:02X}{:02X}{:02X}".format(*rgb)


def _gradient(fill, stops: list[RGB], angle: float) -> None:
    """A linear gradient; ``angle`` in degrees clockwise from left-to-right (DrawingML's convention)."""
    fill.gradient()
    grad = fill._fill._gradFill
    stop_list = grad.find(qn("a:gsLst"))
    for old in list(stop_list):
        stop_list.remove(old)
    for i, rgb in enumerate(stops):
        stop = etree.SubElement(stop_list, qn("a:gs"))
        stop.set("pos", str(round(100000 * i / max(1, len(stops) - 1))))
        etree.SubElement(stop, qn("a:srgbClr")).set("val", _hex(rgb))
    lin = grad.find(qn("a:lin"))
    if lin is None:
        lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", str(round(angle * 60000)))
    lin.set("scaled", "0")


def _soft_shadow(shape) -> None:
    sp_pr = shape._element.spPr
    effects = etree.SubElement(sp_pr, qn("a:effectLst"))
    shadow = etree.SubElement(effects, qn("a:outerShdw"), blurRad="190500", dist="38100", dir="5400000",
                              algn="t", rotWithShape="0")
    colour = etree.SubElement(shadow, qn("a:srgbClr"), val=_hex(JADE_900))
    etree.SubElement(colour, qn("a:alpha"), val="9000")


def _shape(slide, kind, x: float, y: float, w: float, h: float):
    shape = slide.shapes.add_shape(kind, _emu(x), _emu(H - y - h), _emu(w), _emu(h))
    shape.shadow.inherit = False
    shape.line.fill.background()
    return shape


def _fill(slide, s: Fill, shadow: bool) -> None:
    shape = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE if s.radius else MSO_SHAPE.RECTANGLE, s.x, s.y, s.w, s.h)
    if s.radius:
        shape.adjustments[0] = min(0.5, s.radius / max(1e-6, min(s.w, s.h)))
    if s.gradient:
        _gradient(shape.fill, s.gradient, 270 if s.vertical else 0)
    elif s.colour is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*s.colour)
    else:
        shape.fill.background()
    if s.stroke is not None:
        shape.line.color.rgb = RGBColor(*s.stroke)
        shape.line.width = _emu(s.stroke_width)
    if shadow:
        _soft_shadow(shape)


def _disc(slide, s: Disc) -> None:
    shape = _shape(slide, MSO_SHAPE.OVAL, s.cx - s.r, s.cy - s.r, 2 * s.r, 2 * s.r)
    if s.gradient:
        _gradient(shape.fill, s.gradient, 45)
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*s.colour)


def _line(slide, s: Line) -> None:
    size = max(span.size for span in s.spans)
    width = s.width + 6
    x = s.x - (s.width if s.align == "right" else s.width / 2 if s.align == "center" else 0)
    box = slide.shapes.add_textbox(_emu(x), _emu(H - s.y - size * 0.95), _emu(width), _emu(size * 1.25))
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    for span in s.spans:
        run = paragraph.add_run()
        run.text = span.text
        family, bold = PPTX_FONTS[span.font]
        run.font.name = family
        run.font.bold = bold
        run.font.size = Pt(span.size)
        run.font.color.rgb = RGBColor(*span.colour)
        if s.tracking:
            run.font._rPr.set("spc", str(round(s.tracking * 100)))


def _picture(slide, s: Picture) -> None:
    slide.shapes.add_picture(io.BytesIO(s.png), _emu(s.x), _emu(H - s.y - s.h), _emu(s.w), _emu(s.h))


def draw_page(slide, page: Page) -> None:
    _gradient(slide.background.fill, WASH, 45)
    shadow = False
    for shape in page.shapes:
        if isinstance(shape, Shadow):
            shadow = True  # applied to the card drawn next
            continue
        if isinstance(shape, Fill):
            _fill(slide, shape, shadow)
        elif isinstance(shape, Disc):
            _disc(slide, shape)
        elif isinstance(shape, Line):
            _line(slide, shape)
        elif isinstance(shape, Picture):
            _picture(slide, shape)
        shadow = False
    for line in page.hidden:  # adversarial decks: white text a reader cannot see, plus the same in the notes
        _line(slide, line)
    if page.notes:
        slide.notes_slide.notes_text_frame.text = page.notes


def _stable_zip(data: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():  # keeps [Content_Types].xml first
            fixed = zipfile.ZipInfo(entry.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            target.writestr(fixed, source.read(entry.filename))
    return out.getvalue()


def render_slides_pptx(brand: Brand, slides: list[Slide]) -> bytes:
    deck = Presentation()
    deck.slide_width, deck.slide_height = _emu(W), _emu(H)
    blank = deck.slide_layouts[6]
    for number, slide in enumerate(slides, start=1):
        draw_page(deck.slides.add_slide(blank), lay_out(brand, slide, number, len(slides)))
    props = deck.core_properties
    props.title = f"{brand.name} — pitch deck"
    props.author = props.last_modified_by = brand.legal_name
    props.created = props.modified = FIXED_DATE
    props.revision = 1
    buf = io.BytesIO()
    deck.save(buf)
    return _stable_zip(buf.getvalue())


def render_pptx(spec: StartupSpec, slides: list[Slide] | None = None) -> bytes:
    return render_slides_pptx(Brand.from_spec(spec), slides or build_slides(spec))
