"""Document text extraction (Docling → pypdf → pdfplumber) and chunking."""

import asyncio
import logging
import os
import tempfile
import threading
from typing import Dict, List, Optional

from langchain_text_splitters import RecursiveCharacterTextSplitter

from aura.config import DOCLING_ENABLED

logger = logging.getLogger("aura.documents")

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
_MIN_TEXT = 40

_docling_converter = None
_docling_failed = False
_docling_lock = threading.Lock()


def _get_docling():
    """Build Docling and load its PDF models once. Called from a background thread at startup."""
    global _docling_converter, _docling_failed
    if not DOCLING_ENABLED:
        return None
    with _docling_lock:
        if _docling_converter is not None or _docling_failed:
            return _docling_converter
        try:
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions
            from docling.document_converter import DocumentConverter, PdfFormatOption

            # Layout + table structure without OCR: pitch decks carry a text layer, and skipping OCR
            # avoids a large model download and keeps conversion to a few seconds on CPU.
            pdf_options = PdfPipelineOptions(do_ocr=False, do_table_structure=True)
            converter = DocumentConverter(
                format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)}
            )
            converter.initialize_pipeline(InputFormat.PDF)
            _docling_converter = converter
            logger.info("Docling converter ready")
        except Exception as e:
            _docling_failed = True
            logger.warning(f"Docling unavailable, using PDF fallbacks: {e}")
        return _docling_converter


def _extract_docling(path: str) -> str:
    # Never block an upload on Docling's first-time model load; fall back until it is warm.
    if _docling_converter is None:
        return ""
    return _docling_converter.convert(path).document.export_to_markdown()


def _extract_pypdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_pdfplumber(path: str) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return "\n\n".join((page.extract_text() or "") for page in pdf.pages)


def _extract_docx(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_sync(file_bytes: bytes, suffix: str) -> str:
    suffix = (suffix or ".pdf").lower()
    if suffix in (".txt", ".md", ".csv"):
        return file_bytes.decode("utf-8", errors="ignore")

    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(file_bytes)
        path = tmp.name
    try:
        if suffix in (".docx", ".doc"):
            extractors = [("docling", _extract_docling), ("python-docx", _extract_docx)]
        else:
            extractors = [
                ("docling", _extract_docling),
                ("pypdf", _extract_pypdf),
                ("pdfplumber", _extract_pdfplumber),
            ]
        for name, extractor in extractors:
            try:
                text = extractor(path)
            except Exception as e:
                logger.warning(f"{name} extraction failed: {e}")
                continue
            if text and len(text.strip()) >= _MIN_TEXT:
                logger.info(f"Extracted {len(text)} characters with {name}")
                return text
        return ""
    finally:
        os.unlink(path)


async def extract_text(file_bytes: bytes, suffix: str) -> str:
    text = await asyncio.to_thread(_extract_sync, file_bytes, suffix)
    if not text.strip():
        raise ValueError("No readable text found in the uploaded document.")
    return text


def chunk_text(text: str, source: str = "upload") -> List[Dict[str, object]]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return [
        {"text": chunk, "metadata": {"source": source, "chunk_id": i}}
        for i, chunk in enumerate(splitter.split_text(text))
        if chunk.strip()
    ]


def warm_up_docling() -> Optional[str]:
    """Load Docling's models ahead of the first upload. Returns an error message on failure."""
    return None if _get_docling() is not None else "Docling not loaded"

